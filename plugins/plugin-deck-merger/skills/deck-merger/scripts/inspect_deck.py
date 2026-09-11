"""
inspect_deck.py — print a numbered inventory of a .pptx so a human can choose
slides by number.

    python3 inspect_deck.py deck.pptx [deck2.pptx ...]

For each slide it prints: position, HIDDEN marker, layout name, and the first
couple of lines of text. Positions are 1-based and match what PowerPoint shows
in the slide pane, INCLUDING hidden slides — this is the numbering
merge_decks.py expects.

Hidden slides matter: decks maintained as content libraries often hide drafts,
superseded versions and internal-only slides. They are almost never wanted in a
client deck, so they are flagged rather than silently included.
"""
import sys
from pptx import Presentation


def first_lines(slide, n=2):
    out = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for line in shape.text_frame.text.splitlines():
            line = line.strip()
            if line:
                out.append(line[:70])
                break
        if len(out) >= n:
            break
    return out


def inspect(path):
    prs = Presentation(path)
    print(f'\n=== {path}  ({len(prs.slides)} slides)')
    for i, slide in enumerate(prs.slides, 1):
        hidden = 'HIDDEN ' if slide._element.get('show') == '0' else '       '
        text = ' | '.join(first_lines(slide)) or '(no text)'
        print(f'{i:>3}  {hidden}{slide.slide_layout.name[:28]:<30} {text}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    for p in sys.argv[1:]:
        inspect(p)
